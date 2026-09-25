"""Extractor de NOCHES de huéspedes (`Hotel/`, 2013-2016).

Workbooks ACUMULATIVOS (una hoja/noche, copiadas de un libro al siguiente):
la noche se repite entre archivos, muchas hojas sin mes. Cursor (año, mes) en
orden del libro: mes explícito lo fija; sin mes, hereda y avanza si el día
retrocede (31->1). Sin deduplicar entre archivos (eso es 2b): `etl/output/stg_guest_nights.csv`.
"""
from __future__ import annotations

import csv, json, os, re
from dataclasses import asdict, dataclass, field, fields
from datetime import date
from pathlib import Path

from etl.stg_estadias import normalize_payment

ETL_DIR = Path(__file__).resolve().parent.parent
OUT_DIR = ETL_DIR / "output"
HOTEL_ROOT = ETL_DIR.parent / "Hotel"

MONTHS = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}
_MONTH_ABBR = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12,
}
# lunes=0 .. domingo=6, igual que date.weekday()
WEEKDAYS = ["lunes", "martes", "miercoles", "jueves", "viernes", "sabado", "domingo"]
_HOJA_TEMPLATE_RE = re.compile(r"^hoja\s*\d*$")

def _strip_accents_lower(text: str) -> str:
    import unicodedata
    normalized = unicodedata.normalize("NFKD", str(text))
    return "".join(ch for ch in normalized if not unicodedata.combining(ch)).lower()

def _month_from_token(tok: str) -> int | None:
    return MONTHS.get(tok) or _MONTH_ABBR.get(tok[:3])

_DIGIT_RE = re.compile(r"\d{1,2}")
_ALPHA_RE = re.compile(r"[a-z]+")

def _parse_sheet_tokens(sheet_name: str) -> tuple[int | None, int | None, int | None]:
    """Devuelve (día, mes_explícito, weekday_declarado) sin usar cursor."""
    norm = _strip_accents_lower(sheet_name).strip()
    if _HOJA_TEMPLATE_RE.match(norm):
        return None, None, None
    digit_m = _DIGIT_RE.search(norm)
    day = int(digit_m.group(0)) if digit_m else None
    month = None
    weekday = None
    tokens = _ALPHA_RE.findall(norm)
    for tok in tokens:
        if tok in WEEKDAYS:
            weekday = WEEKDAYS.index(tok)
    for tok in tokens:
        if tok in WEEKDAYS:
            continue  # evita que el prefijo de un día de semana matchee un mes (martes ~ mar)
        if month is None:
            month = _month_from_token(tok)
    return day, month, weekday

def sequential_sheet_dates(
    sheet_names: list[str], initial_year: int | None, initial_month: int | None
) -> tuple[list[date | None], list[set[str]]]:
    """Fecha cada hoja en orden con cursor (año, mes); ver docstring del módulo."""
    dated: list[date | None] = []
    flags: list[set[str]] = []
    year, month, last_day = initial_year, initial_month, None

    def unknown():
        dated.append(None)
        flags.append({"date_unknown"})

    for name in sheet_names:
        day, explicit_month, weekday = _parse_sheet_tokens(name)
        if day is None or year is None:
            unknown()
            continue

        if explicit_month is not None:
            if month is not None and explicit_month < month:
                year += 1
            month = explicit_month
        elif month is not None:
            if last_day is not None and day < last_day:
                month += 1
                if month > 12:
                    month, year = 1, year + 1
        else:
            unknown()
            continue

        try:
            d = date(year, month, day)
        except ValueError:
            unknown()
            continue

        f: set[str] = set()
        if weekday is not None and d.weekday() != weekday:
            f.add("weekday_mismatch")
        last_day = day
        dated.append(d)
        flags.append(f)

    return dated, flags

def find_header_row(rows: list[list]) -> int | None:
    for i, row in enumerate(rows):
        joined = " ".join(_strip_accents_lower(c) for c in row if c is not None)
        if "nombre" in joined:
            return i
    return None

_COLUMN_KEYWORDS: list[tuple[str, tuple[str, ...]]] = [
    ("name", ("nombre",)),
    ("company", ("empresa",)),
    ("nationality", ("nacional",)),
    ("pax", ("pax",)),
    ("rate", ("tarifa",)),
    ("payment", ("forma de pago", "forma pago", "pago")),
    ("room", ("habitacion", "hab.", "nro", "no.", "n°", "no", "num")),
]

def detect_columns(header_row: list) -> dict[str, int]:
    cols: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        if cell is None:
            continue
        norm = _strip_accents_lower(cell).strip()
        if not norm:
            continue
        for key, keywords in _COLUMN_KEYWORDS:
            if key in cols:
                continue
            if any(kw in norm for kw in keywords):
                cols[key] = i
                break
    if "room" not in cols:
        cols["room"] = 0
    return cols

def _to_int(v):
    try:
        return int(float(v)) if v is not None else None
    except (ValueError, TypeError):
        return None

def _to_float(v):
    try:
        return float(v) if v is not None else None
    except (ValueError, TypeError):
        return None

def _to_str(v):
    return (str(v).strip() or None) if v is not None else None

@dataclass
class NightObservation:
    night_date: date
    room: int | None
    guest_name: str
    pax: int | None
    rate: float | None
    payment: str | None
    company: str | None
    country: str | None
    source_file: str
    sheet_name: str
    sheet_index: int
    quality_flags: list[str] = field(default_factory=list)

def extract_night_observations(
    rows: list[list], header_idx: int, night_date: date | None, sheet_name: str,
    sheet_index: int, source_file: str, flags: set[str],
) -> list[NightObservation]:
    """Una `NightObservation` por fila ocupada (con nombre) bajo el header."""
    col_map = detect_columns(rows[header_idx])
    obs: list[NightObservation] = []

    def g(row: list, key: str):
        idx = col_map.get(key)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    for row in rows[header_idx + 1:]:
        name = _to_str(g(row, "name"))
        if not name:
            continue
        row_flags = list(flags)
        payment_raw = _to_str(g(row, "payment"))
        obs.append(NightObservation(
            night_date=night_date, room=_to_int(g(row, "room")), guest_name=name,
            pax=_to_int(g(row, "pax")), rate=_to_float(g(row, "rate")),
            payment=normalize_payment(payment_raw) if payment_raw else None,
            company=_to_str(g(row, "company")), country=_to_str(g(row, "nationality")),
            source_file=source_file, sheet_name=sheet_name, sheet_index=sheet_index,
            quality_flags=row_flags,
        ))
    return obs

def capacity_violations(observations: list[NightObservation]) -> list[tuple[str, date, int]]:
    """(source_file, date, count) para fechas con más de 36 habitaciones ocupadas."""
    by_key: dict[tuple[str, date], set[int]] = {}
    for o in observations:
        if o.night_date is None or o.room is None:
            continue
        by_key.setdefault((o.source_file, o.night_date), set()).add(o.room)
    return sorted(
        (src, d, len(rooms)) for (src, d), rooms in by_key.items() if len(rooms) > 36
    )

def _read_xlsx_sheets(path: Path) -> list[tuple[str, list[list]]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        if not hasattr(ws, "iter_rows"):
            continue
        sheets.append((name, [list(row) for row in ws.iter_rows(values_only=True)]))
    return sheets

def _read_xls_sheets(path: Path) -> list[tuple[str, list[list]]]:
    import xlrd
    wb = xlrd.open_workbook(str(path))
    return [(ws.name, [ws.row_values(r) for r in range(ws.nrows)]) for ws in wb.sheets()]

_FILENAME_YEAR_RE = re.compile(r"(20\d\d)")

def process_workbook(path: Path, inventory_year: int | None = None) -> list[NightObservation]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        sheets = _read_xlsx_sheets(path)
    elif suffix == ".xls":
        sheets = _read_xls_sheets(path)
    else:
        return []

    year_m = _FILENAME_YEAR_RE.search(path.name)
    initial_year = int(year_m.group(1)) if year_m else inventory_year

    sheet_names = [name for name, _rows in sheets]
    dated, flag_sets = sequential_sheet_dates(sheet_names, initial_year, None)

    observations: list[NightObservation] = []
    for i, (name, rows) in enumerate(sheets):
        header_idx = find_header_row(rows)
        if header_idx is None:
            continue
        observations.extend(extract_night_observations(
            rows, header_idx, dated[i], name, i, path.name, flag_sets[i],
        ))
    return observations

def _is_guest_register_path(rel_path: str) -> bool:
    norm = _strip_accents_lower(rel_path)
    return "huesp" in norm or "registro de hu" in norm

OBS_COLUMNS = [f.name for f in fields(NightObservation)]

def run() -> None:
    inventory_path = OUT_DIR / "inventory.json"
    if not inventory_path.exists():
        raise SystemExit(f"No existe {inventory_path} — correr primero el extractor (PR1)")

    inventory = json.loads(inventory_path.read_text(encoding="utf-8"))
    candidates = [r for r in inventory if r["is_canonical"] and _is_guest_register_path(r["path"])]

    all_obs: list[NightObservation] = []
    unparsed: list[str] = []
    for rec in candidates:
        path = HOTEL_ROOT / rec["path"]
        if path.suffix.lower() not in (".xlsx", ".xls"):
            unparsed.append(rec["path"])
            continue
        try:
            obs = process_workbook(path, inventory_year=rec.get("year"))
        except Exception as exc:  # noqa: BLE001
            unparsed.append(f"{rec['path']} ({exc})")
            continue
        if not obs:
            unparsed.append(rec["path"])
            continue
        all_obs.extend(obs)

    os.makedirs(OUT_DIR, exist_ok=True)
    csv_path = OUT_DIR / "stg_guest_nights.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(OBS_COLUMNS)
        for o in all_obs:
            d = asdict(o)
            d["quality_flags"] = "|".join(o.quality_flags)
            w.writerow([d[c] for c in OBS_COLUMNS])

    violations = capacity_violations(all_obs)
    dated = [o for o in all_obs if o.night_date is not None]
    by_year: dict[int, int] = {}
    for o in dated:
        by_year[o.night_date.year] = by_year.get(o.night_date.year, 0) + 1
    weekday_mismatch = sum(1 for o in all_obs if "weekday_mismatch" in o.quality_flags)
    pct_unknown = 100 * (len(all_obs) - len(dated)) / len(all_obs) if all_obs else 0

    print(f"OK -> {csv_path}  ({len(all_obs):,} observaciones de noche)")
    print(f"Archivos sin filas parseables: {len(unparsed)}: {unparsed}")
    print(f"% date_unknown: {pct_unknown:.1f}%  weekday_mismatch: {weekday_mismatch}")
    print(f"Noches por año: {dict(sorted(by_year.items()))}")
    print(f"Violaciones de capacidad (>36 hab. mismo día, mismo archivo): {len(violations)}")
    for src, d, n in violations[:20]:
        print(f"  - {src} {d}: {n} habitaciones")

if __name__ == "__main__":
    run()
